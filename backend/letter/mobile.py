"""공문 모바일용 이미지: 같은 내용을 스마트폰 폭(360dp, 3배 = 1080px)에 맞춰 한 줄 배치로 다시 그린 세로 PNG.

A4 공문(PPTX)을 찍은 이미지는 휴대폰에서 확대해야 읽혀서, 표는 항목별 카드로, 가로로 늘어선 지표는 2열 격자로
바꿔 직접 그린다(Pillow). PowerPoint·LibreOffice가 필요 없어 서버와 PC에서 결과가 같다.
- 기업별 값: generator.build_letter_data (PPTX와 같은 값·판정)
- 고정 문구(지원사업 안내·진행 절차·문의처)와 아이콘: 공문 템플릿(template.pptx)에서 읽는다 — 템플릿을 고치면 같이 바뀐다
- 글꼴: Pretendard(SIL OFL, fonts/). '━ ╌ ☎'처럼 글꼴에 없는 기호는 쓰지 않고 선·아이콘으로 그린다
"""
from __future__ import annotations

import io
import math
import re
from functools import lru_cache
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFont
from pptx import Presentation

from ..models import Company
from . import generator as g

S = 3                       # dp -> px 배율
W_DP, MARGIN_DP = 360, 20
W = W_DP * S
FONT_DIR = Path(__file__).with_name("fonts")

NAVY, TEXT, SUB, MUTED = "#16325C", "#1E2430", "#5A6B82", "#6B7A8E"
GREEN, DARK_GREEN, ORANGE = "#1E8C5A", "#1E4D3A", "#B26A00"
LIGHT_BG, LINE, CARD_LINE = "#EEF3F8", "#DDE5EE", "#E2E8F0"
PURPLE = "#7A6FB8"


def u(v: float) -> int:
    return round(v * S)


@lru_cache(maxsize=64)
def font(size_dp: float, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_DIR / f"Pretendard-{'Bold' if bold else 'Regular'}.otf"), u(size_dp))


def _hex(c) -> str:
    return c if str(c).startswith("#") else f"#{c}"


# ---------------------------------------------------------------------------
# 템플릿에서 고정 문구·아이콘 읽기
# ---------------------------------------------------------------------------
Run = tuple[str, bool, str]  # (문구, 굵게, 색)


@lru_cache(maxsize=1)
def _template() -> dict:
    prs = Presentation(str(g.TEMPLATE))
    out: dict = {}
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            out[(si, sh.name)] = sh
    return out


def _paras(si: int, name: str) -> list[list[Run]]:
    sh = _template()[(si, name)]
    return [[(r.text, bool(r.font.bold), _run_color(r)) for r in p.runs] for p in sh.text_frame.paragraphs if p.runs]


def _run_color(r, default=TEXT) -> str:
    try:
        return _hex(str(r.font.color.rgb))
    except Exception:  # 테마 색 등
        return default


def _text(si: int, name: str) -> str:
    return " ".join("".join(t for t, _, _ in p) for p in _paras(si, name))


def _table(si: int, name: str) -> list[list[list[Run]]]:
    table = _template()[(si, name)].table
    return [[[(r.text, bool(r.font.bold), _run_color(r)) for p in c.text_frame.paragraphs for r in p.runs] for c in row.cells]
            for row in table.rows]


@lru_cache(maxsize=8)
def _icon(si: int, name: str, size_dp: float) -> Image.Image:
    img = Image.open(io.BytesIO(_template()[(si, name)].image.blob)).convert("RGBA")
    return img.resize((u(size_dp), round(u(size_dp) * img.height / img.width)), Image.LANCZOS)


@lru_cache(maxsize=4)
def _sprouts(height_px: int) -> Image.Image:
    """배너 그림(그라데이션 위 새싹)에서 새싹만 반투명 흰색으로 뽑는다. 그림 맨 윗줄이 새싹 없는 배경이라
    열마다 그 색을 빼면 새싹 부분만 밝게 남는다."""
    img = Image.open(io.BytesIO(_template()[(1, "Image 2")].image.blob)).convert("RGB")
    bg = img.crop((0, 0, img.width, 1)).resize(img.size)
    # 그림 아래쪽은 배경도 조금 밝아지므로(세로 그라데이션) 배경보다 확실히 밝은 곳만 새싹으로 본다
    diff = ImageChops.subtract(img.convert("L"), bg.convert("L")).point(lambda v: 0 if v < 30 else min(255, (v - 30) * 4))
    edge = round(img.width * 0.01)  # 그림 좌우 가장자리의 밝은 테두리는 뺀다
    diff = diff.crop((edge, 0, img.width - edge, img.height))
    diff = diff.crop(diff.getbbox() or (0, 0, diff.width, diff.height))
    layer = Image.new("RGBA", diff.size, (225, 245, 232, 0))
    layer.putalpha(diff.point(lambda v: v * 0.8))
    return layer.resize((round(layer.width * height_px / layer.height), height_px), Image.LANCZOS)


# ---------------------------------------------------------------------------
# 그리기 도구
# ---------------------------------------------------------------------------
class Canvas:
    def __init__(self):
        self.img = Image.new("RGB", (W, u(6000)), "white")
        self.d = ImageDraw.Draw(self.img)

    def rect(self, x, y, w, h, fill=None, outline=None, width_dp=1.0, radius_dp=10):
        self.d.rounded_rectangle([x, y, x + w, y + h], radius=u(radius_dp), fill=fill, outline=outline,
                                 width=u(width_dp) if outline else 0)

    def hline(self, x, y, w, color=LINE, width_dp=1.0):
        self.d.rectangle([x, y, x + w, y + max(1, u(width_dp)) - 1], fill=color)

    def text(self, x, y, s, size, bold=False, color=TEXT, anchor="la"):
        self.d.text((x, y), s, font=font(size, bold), fill=color, anchor=anchor)

    def paste(self, im: Image.Image, x, y):
        self.img.paste(im, (x, y), im if im.mode == "RGBA" else None)

    def pill(self, x, y, s, size, fill, color, bold=True, pad_dp=10, h_dp=26) -> int:
        """둥근 알약 모양 라벨. 오른쪽 끝 x를 돌려준다."""
        w = round(font(size, bold).getlength(s)) + u(pad_dp) * 2
        self.d.rounded_rectangle([x, y, x + w, y + u(h_dp)], radius=u(h_dp) // 2, fill=fill)
        self.text(x + w // 2, y + u(h_dp) // 2, s, size, bold, color, anchor="mm")
        return x + w


_TOKEN = re.compile(r"[^ ]+ *| +")  # 일반 띄어쓰기에서만 끊는다(U+00A0은 붙임)


def layout(runs: list[Run], width: int, size: float) -> list[list[tuple[str, bool, str, int]]]:
    """어절(띄어쓰기) 단위 줄바꿈. 한 어절이 한 줄보다 길면 글자 단위로 자른다. [(문구, 굵게, 색, x)] 줄 목록."""
    lines, line, x = [], [], 0
    for text, bold, color in runs:
        f = font(size, bold)
        for tok in _TOKEN.findall(text):
            tw = f.getlength(tok.rstrip())
            if x + tw > width and line:
                lines.append(line)
                line, x = [], 0
                if not tok.strip():
                    continue
            if tw > width:  # 너무 긴 어절: 글자 단위
                for ch in tok:
                    cw = f.getlength(ch)
                    if x + cw > width and line:
                        lines.append(line)
                        line, x = [], 0
                    line.append((ch, bold, color, round(x)))
                    x += cw
                continue
            line.append((tok, bold, color, round(x)))
            x += f.getlength(tok)
    if line:
        lines.append(line)
    return lines


def draw_runs(c: Canvas, x, y, width, runs: list[Run], size, line_h=1.6, align="left") -> int:
    """여러 서식이 섞인 문단을 줄바꿈해 그리고 다음 y를 돌려준다."""
    lh = u(size * line_h)
    for line in layout(runs, width, size):
        line_w = 0
        if line:
            t, b, _, lx = line[-1]
            line_w = lx + font(size, b).getlength(t.rstrip())
        off = (width - line_w) / 2 if align == "center" else (width - line_w) if align == "right" else 0
        base = y + lh // 2 + u(size * 0.36)
        for t, b, color, lx in line:
            c.text(x + off + lx, base, t, size, b, color, anchor="ls")
        y += lh
    return y


def text_height(runs: list[Run], width, size, line_h=1.6) -> int:
    return len(layout(runs, width, size)) * u(size * line_h)


def section_header(c: Canvas, y: int, num: str, title: str, caption: str | None) -> int:
    x = u(MARGIN_DP)
    c.rect(x, y, u(28), u(28), fill=NAVY, radius_dp=6)
    c.text(x + u(14), y + u(14), num, 15, True, "white", anchor="mm")
    c.text(x + u(38), y + u(14), title, 19, True, NAVY, anchor="lm")
    y += u(36)
    if caption:
        y = draw_runs(c, x, y, u(W_DP - 2 * MARGIN_DP), [(caption, False, SUB)], 11.5, 1.5)
    return y + u(12)


# ---------------------------------------------------------------------------
# 오각형 그래프
# ---------------------------------------------------------------------------
def radar_image(d: g.LetterData, width_px: int, height_px: int) -> Image.Image:
    k = 2  # 2배로 그린 뒤 줄여 선을 부드럽게
    W2, H2 = width_px * k, height_px * k
    im = Image.new("RGBA", (W2, H2), (255, 255, 255, 0))
    dr = ImageDraw.Draw(im)
    cx, cy, r = W2 // 2, H2 // 2 + u(6) * k, u(78) * k
    n = len(d.radar)

    def pt(i, frac):
        ang = -math.pi / 2 + 2 * math.pi * i / n
        return cx + r * frac * math.cos(ang), cy + r * frac * math.sin(ang)

    for frac in (0.25, 0.5, 0.75, 1.0):
        dr.polygon([pt(i, frac) for i in range(n)], outline="#D5DCE5", width=k * 2)
    for i in range(n):
        dr.line([(cx, cy), pt(i, 1.0)], fill="#E3E8EF", width=k * 2)
    has_data = any(a.score for a in d.radar)
    comp = [pt(i, a.score / 100) for i, a in enumerate(d.radar)]
    if has_data:  # 반투명 색칠은 별도 층에 그려 합성 (같은 층에 그리면 아래 선을 덮어쓴다)
        fill = Image.new("RGBA", im.size, (0, 0, 0, 0))
        ImageDraw.Draw(fill).polygon(comp, fill=(30, 140, 90, 40))
        im = Image.alpha_composite(im, fill)
        dr = ImageDraw.Draw(im)
    avg = [pt(i, g.INDUSTRY_AVG_SCORE / 100) for i in range(n)]
    for a, b in zip(avg, avg[1:] + avg[:1]):  # 업종평균: 점선
        steps = 9
        for s in range(0, steps, 2):
            p0 = (a[0] + (b[0] - a[0]) * s / steps, a[1] + (b[1] - a[1]) * s / steps)
            p1 = (a[0] + (b[0] - a[0]) * (s + 1) / steps, a[1] + (b[1] - a[1]) * (s + 1) / steps)
            dr.line([p0, p1], fill=PURPLE, width=u(2) * k)
    for p in avg:
        dr.ellipse([p[0] - u(3) * k, p[1] - u(3) * k, p[0] + u(3) * k, p[1] + u(3) * k], fill=PURPLE)
    if has_data:
        dr.line(comp + comp[:1], fill=GREEN, width=u(3) * k, joint="curve")
        for p in comp:
            dr.ellipse([p[0] - u(4) * k, p[1] - u(4) * k, p[0] + u(4) * k, p[1] + u(4) * k], fill=GREEN)
    im = im.resize((width_px, height_px), Image.LANCZOS)

    # 축 이름(굵게)과 등급(괄호) — 원래 크기에서 글자 그리기
    dr = ImageDraw.Draw(im)
    cx, cy, r = width_px // 2, height_px // 2 + u(6), u(78)
    for i, a in enumerate(d.radar):
        ang = -math.pi / 2 + 2 * math.pi * i / n
        lx, ly = cx + (r + u(22)) * math.cos(ang), cy + (r + u(22)) * math.sin(ang)
        h_anchor = "m" if abs(math.cos(ang)) < 0.3 else ("l" if math.cos(ang) > 0 else "r")
        if i == 0:
            ly -= u(10)
        dr.text((lx, ly - u(8)), a.label, font=font(12.5, True), fill=NAVY, anchor=h_anchor + "m")
        dr.text((lx, ly + u(8)), f"({a.grade or '자료없음'})", font=font(11.5), fill=SUB, anchor=h_anchor + "m")
    if not any(a.score for a in d.radar):
        msg = "재무진단 자료 없음"
        half = font(13, True).getlength(msg) / 2 + u(10)
        dr.rounded_rectangle([cx - half, cy - u(14), cx + half, cy + u(14)], radius=u(8), fill="white", outline=LINE, width=u(1))
        dr.text((cx, cy), msg, font=font(13, True), fill=MUTED, anchor="mm")
    return im


# ---------------------------------------------------------------------------
# 본문
# ---------------------------------------------------------------------------
def render_mobile_png(c_: Company | g.LetterData, *, doc_no: str = "", issued=None) -> bytes:
    d = c_ if isinstance(c_, g.LetterData) else g.build_letter_data(c_, doc_no=doc_no, issued=issued)
    c = Canvas()
    X, CW = u(MARGIN_DP), u(W_DP - 2 * MARGIN_DP)
    y = u(26)

    # ----- 머리: 로고 · 슬로건
    c.paste(_icon(1, "Image 0", 30), X, y)
    c.text(X + u(38), y + u(1), "FARMSCOPE", 21, True, NAVY)
    c.text(X + u(39), y + u(27), "A G R I   B U S I N E S S   P A R T N E R", 7, True, MUTED)
    slogan = _paras(1, "Text 2")
    c.text(X + CW, y + u(4), "".join(t for t, _, _ in slogan[0]), 11.5, True, DARK_GREEN, anchor="ra")
    if len(slogan) > 1:
        c.text(X + CW, y + u(21), "".join(t for t, _, _ in slogan[1]), 11.5, True, DARK_GREEN, anchor="ra")
    y += u(48)
    c.hline(X, y, CW, NAVY, 2)
    y += u(16)

    # ----- 문서 정보
    LW = u(62)
    title = _text(1, "Text 11")
    for label, runs, size in (
        ("문서번호", [(d.doc_no, False, TEXT)], 13.5),
        ("시행일자", [(d.issued_text, False, TEXT)], 13.5),
        ("수    신", [(d.recipient_name, True, TEXT), (d.recipient_suffix, False, TEXT)], 13.5),
        ("제    목", [(title, True, NAVY)], 14.5),
    ):
        c.text(X, y + u(size * 0.8), label.replace("    ", "  "), 13, True, SUB, anchor="ls")
        y = draw_runs(c, X + LW, y - u(size * 0.3), CW - LW, runs, size, 1.55) + u(4)
    y += u(6)
    c.hline(X, y, CW, LINE)
    y += u(16)

    # ----- 인사말
    for para in _paras(1, "Text 13"):
        y = draw_runs(c, X, y, CW, para, 14.5, 1.7) + u(6)
    y += u(10)

    # ----- 약속 상자 (악수 아이콘)
    promise = _paras(1, "Text 15")
    PAD, IW = u(16), u(40)
    sizes = [16, 16, 12.5]
    inner = CW - PAD * 2 - IW - u(12)
    h = PAD * 2 + sum(text_height(p, inner, s, 1.5) for p, s in zip(promise, sizes)) + u(4)
    c.rect(X, y, CW, h, fill="#E9F4EC", radius_dp=12)
    c.paste(_icon(1, "Image 1", 40), X + PAD, y + (h - u(40)) // 2)
    ty = y + PAD
    for p, s in zip(promise, sizes):
        ty = draw_runs(c, X + PAD + IW + u(12), ty, inner, p, s, 1.5) + (u(4) if s != 16 else 0)
    y += h + u(32)

    # ===== 1. 지원사업 주요내용
    y = section_header(c, y, "1", _text(1, "Text 18"), _text(1, "Text 19"))
    # 초록 배너 (그라데이션 + 새싹 그림)
    BH = u(112)
    grad = Image.new("RGB", (256, 1))
    for i in range(256):
        t = i / 255
        grad.putpixel((i, 0), (round(20 + (90 - 20) * t), round(70 + (165 - 70) * t), round(45 + (120 - 45) * t)))
    grad = grad.resize((CW, BH))
    mask = Image.new("L", (CW, BH), 0)
    ImageDraw.Draw(mask).rounded_rectangle([0, 0, CW - 1, BH - 1], radius=u(12), fill=255)
    sprouts = _sprouts(u(34))
    grad.paste(sprouts, (CW - sprouts.width - u(6), BH - sprouts.height), sprouts)
    c.img.paste(grad, (X, y), mask)
    ty = draw_runs(c, X + u(18), y + u(18), CW - u(36), [(_text(1, "Text 20"), True, "#FFFFFF")], 19, 1.35)
    draw_runs(c, X + u(18), ty + u(4), CW - u(36), [(_text(1, "Text 21"), False, "#E8F5EC")], 12, 1.45)
    y += BH + u(12)

    # 지표 카드 2x2 (융자금리·법인 한도·상환기간·신청 마감)
    GAP = u(10)
    CWH = (CW - GAP) // 2
    cards = [("Text 23", "Text 24", "Text 25"), ("Text 27", "Text 28", "Text 29"),
             ("Text 31", "Text 32", "Text 33"), ("Text 35", "Text 36", "Text 37")]
    for i, (lab, val, sub) in enumerate(cards):
        cx_, cy_ = X + (i % 2) * (CWH + GAP), y + (i // 2) * (u(100) + GAP)
        c.rect(cx_, cy_, CWH, u(100), fill="white", outline="#CFE6D8", width_dp=1.5, radius_dp=10)
        c.text(cx_ + CWH // 2, cy_ + u(20), _text(1, lab), 12.5, False, SUB, anchor="mm")
        v = _paras(1, val)[0]
        vsize = 22 if font(22, True).getlength("".join(t for t, _, _ in v)) < CWH - u(16) else 18
        c.text(cx_ + CWH // 2, cy_ + u(50), "".join(t for t, _, _ in v), vsize, True, v[0][2], anchor="mm")
        draw_runs(c, cx_ + u(8), cy_ + u(68), CWH - u(16), [(_text(1, sub), False, MUTED)], 11, 1.35, align="center")
    y += 2 * u(100) + GAP + u(14)

    # 안내 표 (지원대상·지원용도·상환조건·신청방법): 라벨 칸 + 내용
    LW2 = u(74)
    for row in _table(1, "Table 0"):
        label, value = row[0], row[1]
        th = max(text_height(value, CW - LW2 - u(24), 13.5, 1.55), u(22)) + u(24)
        c.d.rectangle([X, y, X + LW2, y + th], fill=LIGHT_BG)
        c.text(X + u(12), y + th // 2, "".join(t for t, _, _ in label), 13.5, True, NAVY, anchor="lm")
        draw_runs(c, X + LW2 + u(12), y + u(12), CW - LW2 - u(24), value, 13.5, 1.55)
        y += th
        c.hline(X, y, CW, LINE)
    y += u(16)

    # 주의 문구 (! 아이콘)
    c.paste(_icon(1, "Image 3", 20), X, y + u(2))
    for para in _paras(1, "Text 38"):
        y = draw_runs(c, X + u(30), y, CW - u(30), para, 13, 1.6) + u(4)
    y += u(34)

    # ===== 2. 재무진단 요약
    y = section_header(c, y, "2", _text(2, "Text 6"), d.basis_text)
    rows = [("주요사업", d.industry), ("설립연도", d.founded_text), ("소재지", d.location)]
    PAD = u(18)
    inner = CW - PAD * 2 - u(66)
    h = PAD * 2 + u(34) + sum(text_height([(v, False, TEXT)], inner, 13.5, 1.5) + u(6) for _, v in rows)
    c.rect(X, y, CW, h, fill=LIGHT_BG, radius_dp=12)
    ty = draw_runs(c, X + PAD, y + PAD, CW - PAD * 2, [(d.name, True, NAVY)], 19, 1.3) + u(8)
    for lab, v in rows:
        c.text(X + PAD, ty + u(10), lab, 13, True, SUB, anchor="lm")
        ty = draw_runs(c, X + PAD + u(66), ty, inner, [(v, False, TEXT)], 13.5, 1.5) + u(6)
    y += h + u(10)

    # 오각형 그래프 + 범례
    RH = u(260)
    c.paste(radar_image(d, CW, RH), X, y)
    y += RH
    lx = X
    c.d.line([lx, y + u(9), lx + u(22), y + u(9)], fill=GREEN, width=u(3))
    lx += u(28)
    c.text(lx, y + u(9), d.name, 12.5, True, TEXT, anchor="lm")
    lx += round(font(12.5, True).getlength(d.name)) + u(18)
    for s in range(3):
        c.d.line([lx + u(s * 8), y + u(9), lx + u(s * 8 + 5), y + u(9)], fill=PURPLE, width=u(3))
    c.text(lx + u(28), y + u(9), "동종업종 평균", 12.5, False, TEXT, anchor="lm")
    y += u(24)
    scale = " · ".join(f"{k} {v}".replace(" ", "\u00a0") for k, v in g.GRADE_SCORE.items())  # '낮음 20'이 끊기지 않게
    y = draw_runs(c, X, y, CW, [("KODATA 등급 환산  ", True, MUTED), (scale, False, MUTED)], 11.5, 1.5) + u(14)

    # 재무 지표 카드: 2열 (매출액·영업이익 / 부채비율·유동비율 / 신용등급 전체 폭)
    KH = u(104)
    for i, k in enumerate(d.kpis):
        full = i == 4
        kx = X if full or i % 2 == 0 else X + CWH + GAP
        kw = CW if full else CWH
        ky = y + (i // 2) * (KH + GAP)
        c.rect(kx, ky, kw, KH, fill="white", outline=CARD_LINE, width_dp=1.2, radius_dp=10)
        c.d.rounded_rectangle([kx + u(1), ky + u(1), kx + kw - u(1), ky + u(30)], radius=u(9), fill=LIGHT_BG,
                              corners=(True, True, False, False))
        c.text(kx + kw // 2, ky + u(16), k.label, 13, True, NAVY, anchor="mm")
        val_w = font(24, True).getlength(k.value)
        unit = f" {k.unit}" if k.unit and k.value != "-" else ""
        unit_w = font(12.5, True).getlength(unit)
        vx = kx + (kw - val_w - unit_w) / 2
        c.text(vx, ky + u(68), k.value, 24, True, _hex(k.color), anchor="ls")
        if unit:
            c.text(vx + val_w, ky + u(68), unit, 12.5, True, _hex(k.color), anchor="ls")
        c.text(kx + kw // 2, ky + u(88), k.sub, 11.5, False, MUTED, anchor="mm")
    y += 3 * KH + 2 * GAP + u(34)

    # ===== 3. 지원조건 점검 결과
    y = section_header(c, y, "3", _text(2, "Text 44"), None) - u(4)
    px = X
    for key, sym in (("충족", "✓"), ("확인 필요", "?"), ("보완 필요", "!")):
        fill, color, _ = g.JUDGE_STYLE[key]
        px = c.pill(px, y, f"{sym} {key} {d.counts[key]}", 12.5, _hex(fill), _hex(color)) + u(8)
    y += u(26) + u(14)
    for ch in d.checks:
        PAD = u(14)
        inner = CW - PAD * 2
        h = PAD * 2 + u(26) + u(6) + text_height([(ch.status, False, TEXT)], inner, 13.5, 1.5) + u(4) \
            + text_height([("다음 단계  ", True, SUB), (ch.next_step, False, SUB)], inner, 12.5, 1.5)
        fill, color, label = g.JUDGE_STYLE[ch.judgment]
        c.rect(X, y, CW, h, fill="white", outline=CARD_LINE, width_dp=1.2, radius_dp=10)
        c.d.rectangle([X, y + u(10), X + u(4), y + h - u(10)], fill=_hex(color))  # 판정 색 띠
        c.text(X + PAD, y + PAD + u(13), ch.label, 15, True, NAVY, anchor="lm")
        pw = round(font(12.5, True).getlength(label)) + u(20)
        c.pill(X + CW - PAD - pw, y + PAD, label, 12.5, _hex(fill), _hex(color))
        ty = draw_runs(c, X + PAD, y + PAD + u(32), inner, [(ch.status, False, TEXT)], 13.5, 1.5) + u(4)
        draw_runs(c, X + PAD, ty, inner, [("다음 단계  ", True, SUB), (ch.next_step, False, SUB)], 12.5, 1.5)
        y += h + u(8)
    y += u(6)

    # 종합 의견
    styles = {"normal": (False, TEXT), "bold": (True, TEXT), "green": (True, GREEN)}
    vruns = [(t, *styles[s]) for t, s in d.verdict]
    PAD = u(16)
    h = PAD * 2 + u(28) + u(10) + text_height(vruns, CW - PAD * 2, 14.5, 1.65)
    c.rect(X, y, CW, h, fill="#F2F8F4", outline="#A9D5B9", width_dp=1.5, radius_dp=12)
    c.pill(X + PAD, y + PAD, "종합 의견", 13, DARK_GREEN, "#FFFFFF", h_dp=28)
    draw_runs(c, X + PAD, y + PAD + u(38), CW - PAD * 2, vruns, 14.5, 1.65)
    y += h + u(34)

    # ===== 4. 진행 절차 및 상담 신청
    y = section_header(c, y, "4", _text(2, "Text 57"), _text(2, "Text 58"))
    PAD = u(16)
    inner = CW - PAD * 2
    head = [(_text(2, "Text 62"), True, "#FFFFFF")]
    steps = [("Text 65", "Text 66", "Text 67", "Text 68"), ("Text 72", "Text 73", "Text 74", "Text 75"),
             ("Text 79", "Text 80", "Text 81", "Text 82"), ("Text 86", "Text 87", "Text 88", "Text 89")]
    SW = (inner - GAP) // 2
    SH = u(96)
    contact = _paras(2, "Text 91")[0]
    top = y
    box_h = PAD + u(26) + u(10) + text_height(head, inner, 15, 1.45) + u(12) + 2 * SH + GAP + u(18) + u(1) + u(16) + u(96) + PAD
    c.rect(X, y, CW, box_h, fill=NAVY, radius_dp=14)
    y += PAD
    c.pill(X + PAD, y, "ONE-STOP", 11, "#8FE0B4", NAVY, pad_dp=10, h_dp=26)
    y += u(36)
    y = draw_runs(c, X + PAD, y, inner, head, 15, 1.45) + u(12)
    for i, (num, tit, desc, when) in enumerate(steps):
        sx, sy = X + PAD + (i % 2) * (SW + GAP), y + (i // 2) * (SH + GAP)
        c.rect(sx, sy, SW, SH, fill="white", radius_dp=10)
        c.rect(sx + u(10), sy + u(12), u(26), u(20), fill=NAVY, radius_dp=4)
        c.text(sx + u(23), sy + u(22), _text(2, num), 11, True, "white", anchor="mm")
        c.text(sx + u(42), sy + u(22), _text(2, tit), 14, True, NAVY, anchor="lm")
        draw_runs(c, sx + u(10), sy + u(40), SW - u(20), [(_text(2, desc), False, "#3A4555")], 12, 1.4)
        c.text(sx + u(10), sy + u(80), _text(2, when), 12.5, True, ORANGE, anchor="lm")
    y += 2 * SH + GAP + u(18)
    c.hline(X + PAD, y, inner, "#3B5378")
    y += u(16)
    # 문의처: 템플릿 문구(문의 / ☎ 번호 / 휴대폰 번호 / 담당자)를 줄 단위로
    c.paste(_icon(2, "Image 1", 20), X + PAD, y)
    c.text(X + PAD + u(28), y + u(10), contact[0][0].strip(), 15, True, "#FFFFFF", anchor="lm")
    texts = [t.strip() for t, _, _ in contact[1:]]
    if len(texts) >= 5:
        lines = [[("전화  ", False, "#DCE5F0"), (texts[1], True, "#FFFFFF")],
                 [(texts[2] + "  ", False, "#DCE5F0"), (texts[3], True, "#FFFFFF")],
                 [(texts[4], False, "#DCE5F0")]]
    else:
        lines = [[(" ".join(texts), False, "#DCE5F0")]]
    ty = y + u(26)
    for ln in lines:
        ty = draw_runs(c, X + PAD, ty, inner, ln, 13.5 if ln[-1][1] else 12.5, 1.5)
    y = top + box_h + u(22)

    # ----- 하단
    c.hline(X, y, CW, LINE)
    y += u(12)
    y = draw_runs(c, X, y, CW, [(_text(2, "Text 93"), False, MUTED)], 11.5, 1.5) + u(8)
    c.text(X, y + u(10), d.issued_text, 12, False, MUTED, anchor="lm")
    c.text(X + CW, y + u(10), "FARMSCOPE", 15, True, NAVY, anchor="rm")
    y += u(20) + u(28)

    out = c.img.crop((0, 0, W, y))
    buf = io.BytesIO()
    out.save(buf, "PNG", compress_level=6)
    return buf.getvalue()
