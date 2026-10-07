"""영업 관리 '공문 다운로드': 공문 템플릿(template.pptx)의 기업정보 칸을 기업 데이터로 채운다.

template.pptx는 사용자가 준 공문_샘플.pptx 그대로다. 도형은 이름('Text 9' 등)으로 찾으므로 템플릿을 PowerPoint에서
고쳤다면 tests/test_letter.py를 돌려 칸을 여전히 찾는지 확인한다. 지원사업 안내(1쪽 본문, 2쪽 4번 진행 절차)는
모든 기업에 같은 문구라 템플릿 그대로 둔다.

LLM 없이 규칙으로만 채운다: 수치는 KODATA 기업보고서에서 추출한 값, 판정은 아래 _judge_* 규칙, 종합 의견은
판정 결과에 따라 고른 문장틀이다. 데이터로 알 수 없는 항목(담보·투자계획·보조사업 중복)은 '확인 필요'로 둔다.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Pt

from ..models import Company

TEMPLATE = Path(__file__).with_name("template.pptx")

# 템플릿 색 (샘플에서 그대로 가져옴)
NAVY, GREEN, ORANGE = "16325C", "1E8C5A", "B26A00"
JUDGE_STYLE = {  # 판정 -> (표 칸 배경, 글자색, 표시 문구)
    "충족": ("DDF0E4", "17603D", "✓ 충족"),
    "확인 필요": ("FCEFD6", "8A5300", "? 확인 필요"),
    "보완 필요": ("FBE1DE", "A3281C", "! 보완 필요"),
}

# 재무진단 등급 -> 오각형 그래프 점수 (사용자 확정 2026-10-07: 샘플의 양호 70·보통 50·보통 이하 40에 우수·낮음 추가)
GRADE_SCORE = {"우수": 90, "양호": 70, "보통": 50, "보통 이하": 40, "낮음": 20}
INDUSTRY_AVG_SCORE = 50
RADAR_AXES = [("growth", "성장성"), ("profitability", "수익성"), ("financial_structure", "재무구조"),
              ("debt_repayment", "부채상환능력"), ("activity", "활동성")]

# KODATA 신용등급 (좋은 순)
CREDIT_GRADES = ["aaa", "aa+", "aa", "aa-", "a+", "a", "a-", "bbb+", "bbb", "bbb-", "bb+", "bb", "bb-",
                 "b+", "b", "b-", "ccc+", "ccc", "ccc-", "cc", "c", "d"]

TARGET_SIDO = "경상북도"  # 템플릿의 지원사업(2027년도 경상북도 농어촌진흥기금) 대상 지역

SIDO = {
    "서울": "서울특별시", "부산": "부산광역시", "대구": "대구광역시", "인천": "인천광역시", "광주": "광주광역시",
    "대전": "대전광역시", "울산": "울산광역시", "세종": "세종특별자치시", "경기": "경기도", "강원": "강원특별자치도",
    "충북": "충청북도", "충남": "충청남도", "전북": "전북특별자치도", "전남": "전라남도", "경북": "경상북도",
    "경남": "경상남도", "제주": "제주특별자치도",
}
_SIDO_FULL = {**{v: v for v in SIDO.values()}, "강원도": "강원특별자치도", "전라북도": "전북특별자치도",
              "제주도": "제주특별자치도", "서울시": "서울특별시"}

MINUS = "−"  # 샘플의 음수 표기(−)
PAREN = re.compile(r"[(][^)]*[)]")  # 괄호 내용


# ---------------------------------------------------------------------------
# 데이터 가공
# ---------------------------------------------------------------------------
@dataclass
class Address:
    sido: str | None = None
    sigungu: str | None = None
    town: str | None = None  # 읍·면·동

    @property
    def short(self) -> str:
        return " ".join(p for p in (self.sido, self.sigungu, self.town) if p) or "-"


def split_address(address: str | None) -> Address:
    """'경북의성군안계면…(○○리)'처럼 띄어쓰기 없는 주소에서 시도·시군구·읍면동을 뽑는다."""
    s = re.sub(r"\s+", "", address or "")
    out = Address()
    for full in sorted(_SIDO_FULL, key=len, reverse=True):
        if s.startswith(full):
            out.sido, s = _SIDO_FULL[full], s[len(full):]
            break
    else:
        if s[:2] in SIDO:
            out.sido, s = SIDO[s[:2]], s[2:]
    m = re.match(r"([가-힣]+?[시군구])", s)
    if m and out.sido != "세종특별자치시":
        out.sigungu, s = m.group(1), s[m.end():]
        gu = re.match(r"[가-힣]{1,3}구(?![로길])", s)  # 포항시남구 등 일반구
        if out.sigungu.endswith("시") and gu:
            out.sigungu, s = f"{out.sigungu} {gu.group(0)}", s[gu.end():]
    m = re.match(r"([가-힣]{1,4}?[읍면])(?![로길])", s)
    if m:
        out.town = m.group(1)
    else:
        m = re.search(r"\(([가-힣0-9]+?[동가])[,)]", s) or re.match(r"([가-힣0-9]{1,5}?동)(?![로길])", s)
        if m:
            out.town = m.group(1)
    return out


def short_name(name: str) -> str:
    """'농업회사법인 테스트종묘(주)' -> '테스트종묘(주)'. 법인 형태 문구만 지운다."""
    s = re.sub(r"\s+", "", name or "")
    out = re.sub(r"농업회사법인|영농조합법인|영어조합법인|어업회사법인", "", s)
    return out or s


def spaced_name(name: str) -> str:
    """수신란 표기: 앞에 붙은 법인 형태 뒤를 띄운다('농업회사법인테스트종묘(주)' -> '농업회사법인 테스트종묘(주)')."""
    return re.sub(r"^(\(주\))?(농업회사법인|영농조합법인|영어조합법인|어업회사법인)(?=\S)", lambda m: m.group(0) + " ", (name or "").strip())


def org_kind(c: Company) -> str:
    name = c.company_name or ""
    for kind in ("농업회사법인", "영농조합법인", "영어조합법인", "어업회사법인"):
        if kind in name.replace(" ", ""):
            return kind
    return "개인사업자" if c.company_type == "개인사업자" else "법인"


def pretty_industry(name: str | None) -> str:
    if not name:
        return "-"
    s = re.sub(r"\s*및\s*", " 및 ", name)
    return re.sub(r",\s*", ", ", s).strip()


def _latest(series: dict | None) -> tuple[str | None, float | None]:
    nums = {y: v for y, v in (series or {}).items() if y.isdigit() and isinstance(v, (int, float))}
    if not nums:
        return None, None
    y = max(nums)
    return y, nums[y]


def _prev(series: dict | None, year: str | None) -> float | None:
    if not year:
        return None
    v = (series or {}).get(str(int(year) - 1))
    return v if isinstance(v, (int, float)) else None


def indicator(c: Company, axis: str, name: str) -> tuple[float | None, float | None]:
    """재무진단 지표 (기업 값, 업종평균). 진단표에 없으면 비율표의 최근값을 쓴다(업종평균은 없음)."""
    for ind in (c.diagnosis_details.get(axis) or {}).get("indicators", []):
        if ind.get("name") == name:
            comp, avg = ind.get("company"), ind.get("industry_avg")
            return (comp if isinstance(comp, (int, float)) else None, avg if isinstance(avg, (int, float)) else None)
    _, v = _latest(c.ratio_summary.get(name))
    if v is None:
        for table in c.ratio_detail.values():
            if name in table:
                _, v = _latest(table[name])
                break
    return v, None


def growth_rate(c: Company, name: str, cur: float | None, prev: float | None) -> float | None:
    """전년 대비 증감률(%). KODATA가 계산한 증가율을 우선 쓰고, 없으면 직접 계산(전년이 0 이하면 없음)."""
    v, _ = indicator(c, "growth", name)
    if v is not None:
        return v
    if cur is None or prev is None or prev <= 0:
        return None
    return (cur - prev) / prev * 100


def grade_rank(grade: str | None) -> int | None:
    g = (grade or "").strip().lower()
    return CREDIT_GRADES.index(g) if g in CREDIT_GRADES else None


def fmt_num(v: float, digits: int) -> str:
    s = f"{v:,.{digits}f}"
    return s.replace("-", MINUS)


def fmt_eok(million: float) -> str:
    """백만원 -> 억원. 10억 미만은 소수 둘째 자리까지(샘플: 22.5억, 0.55억, 2.73억)."""
    eok = million / 100
    return fmt_num(eok, 1 if abs(eok) >= 10 else 2)


def fmt_signed_pct(v: float) -> str:
    return ("+" if v > 0 else "") + fmt_num(v, 1) + "%"


def norm_grade(g: str | None) -> str | None:
    if not g:
        return None
    return "보통 이하" if g.replace(" ", "") == "보통이하" else g.strip()


# ---------------------------------------------------------------------------
# 지원조건 판정 (사용자 확정 2026-10-07: 수치로 판정 가능한 항목만 자동 판정, 나머지는 '확인 필요')
# ---------------------------------------------------------------------------
@dataclass
class Check:
    label: str
    status: str
    judgment: str
    next_step: str


def _judge_basic(c: Company, addr: Address) -> Check:
    status = f"{addr.sigungu or addr.sido or '소재지 미상'} 소재 {org_kind(c)}"
    if addr.sido == TARGET_SIDO:
        return Check("기본 자격", status, "충족", "대표자 주소(도내)·경영체 등록 확인")
    if addr.sido is None:
        return Check("기본 자격", status, "확인 필요", "사업장 소재지(도내) 확인")
    return Check("기본 자격", status, "보완 필요", "도내 사업장 소재 요건 확인")


def _judge_finance(c: Company, debt: float | None, debt_avg: float | None, equity: float | None) -> Check:
    grade = (c.credit_grade or "").strip()
    rank = grade_rank(grade)
    parts = [f"부채비율 {fmt_num(debt, 0)}%" + (f" (업종 {fmt_num(debt_avg, 0)}%)" if debt_avg is not None else "")
             if debt is not None else "부채비율 자료없음"]
    if grade:
        parts.append(grade)
    status = " · ".join(parts)
    if (equity is not None and equity < 0) or (debt is not None and debt > 400) or (rank is not None and rank >= grade_rank("ccc+")):
        return Check("재무·신용", status, "보완 필요", "재무구조 개선 방안 마련")
    debt_ok = debt is not None and debt <= (debt_avg if debt_avg is not None else 200)
    if debt_ok and rank is not None and rank <= grade_rank("bb-"):
        return Check("재무·신용", status, "충족", "농협 신용조사의견서 발급")
    return Check("재무·신용", status, "확인 필요", "부채 구성·신용등급 사유 확인")


def _judge_liquidity(current: float | None, icr: float | None) -> Check:
    parts = [f"유동비율 {fmt_num(current, 0)}%" if current is not None else "유동비율 자료없음"]
    if icr is not None:
        parts.append(f"이자보상 {fmt_num(icr, 1)}배")
    status = " · ".join(parts)
    if current is None:
        return Check("유동성·상환력", status, "확인 필요", "유동성 자료 확인")
    if current < 100 or (icr is not None and icr < 1):
        return Check("유동성·상환력", status, "보완 필요", "선대출 승인·상환계획 수립")
    return Check("유동성·상환력", status, "충족", "거치기간 활용 상환계획 수립")


def build_checks(c: Company, addr: Address, debt, debt_avg, current, icr, equity) -> list[Check]:
    collateral = "재무상 여력 양호 (추정)" if (equity or 0) > 0 and debt is not None and debt <= 200 else "재무상 여력 제한적 (추정)"
    place = (addr.town or addr.sigungu or "").split(" ")[-1]
    place = place[:-1] if len(place) > 1 else place
    return [
        _judge_basic(c, addr),
        _judge_finance(c, debt, debt_avg, equity),
        Check("담보 여력", collateral, "확인 필요", "보유 부동산·기존 담보 확인"),
        Check("투자계획", f"{place} 사업장 시설 확충 검토" if place else "사업장 시설 확충 검토", "확인 필요", "스마트팜 해당 시 한도 10억원"),
        Check("보조사업 중복", "병행 여부 미확인", "확인 필요", "기금·보조 자금 분리 설계"),
        _judge_liquidity(current, icr),
    ]


def verdict_segments(checks: list[Check]) -> list[tuple[str, str]]:
    """종합 의견 문장 (문구, 서식: normal|bold|green)."""
    basic, finance, liquidity = checks[0].judgment, checks[1].judgment, checks[5].judgment
    if basic != "충족":
        head = [("도내 사업장 요건 ", "normal"), ("확인이 우선", "bold"), ("입니다. 요건 충족 시 ", "normal")]
    elif finance == "충족":
        head = [("기본 자격과 재무구조는 ", "normal"), ("충족", "bold"), ("합니다. 투자계획이 확정되면 ", "normal")]
    else:
        head = [("기본 자격은 충족하나 재무·신용 ", "normal"),
                ("보완이 필요" if finance == "보완 필요" else "확인이 필요", "bold"), ("합니다. 투자계획 확정 시 ", "normal")]
    tail = {"보완 필요": " 신청을 검토할 수 있으며, 유동성을 보완하는 자금계획이 함께 필요합니다.",
            "충족": " 신청을 검토할 수 있으며, 단기 상환 여력도 양호한 편입니다.",
            "확인 필요": " 신청을 검토할 수 있으며, 유동성 자료를 함께 확인해야 합니다."}[liquidity]
    if basic != "충족" or finance != "충족":  # 앞 문장이 길어지면 뒤를 줄여 두 줄 안에 맞춘다
        tail = {"보완 필요": " 신청을 검토할 수 있고, 유동성 보완 계획이 필요합니다.",
                "충족": " 신청을 검토할 수 있고, 단기 상환 여력은 양호합니다.",
                "확인 필요": " 신청을 검토할 수 있고, 유동성 자료 확인이 필요합니다."}[liquidity]
    return head + [("일반 시설 5억원", "green"), (" 또는 ", "normal"), ("스마트팜 최대 10억원", "green"), (tail, "normal")]


# ---------------------------------------------------------------------------
# PPTX 채우기
# ---------------------------------------------------------------------------
def _shape(slide, name: str):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(f"공문 템플릿에서 '{name}' 도형을 찾지 못했습니다.")


def _runs(shape) -> list:
    return [r for p in shape.text_frame.paragraphs for r in p.runs]


def _set_text(shape, *texts: str) -> None:
    runs = _runs(shape)
    for run, text in zip(runs, texts):
        run.text = text


def _text_width(text: str, size: float) -> float:
    """대략의 글자 폭(pt): 한글 1em, 나머지 0.55em."""
    return sum(size if ord(ch) > 0x2E80 else size * 0.55 for ch in text)


def _fit(runs: list, text: str, width_pt: float, min_size: float = 8.0) -> None:
    """한 줄에 들어가도록 글자 크기를 줄인다(최소 min_size)."""
    size = runs[0].font.size.pt
    new = size
    while new > min_size and _text_width(text, new) > width_pt:
        new -= 0.2
    if new < size:
        for r in runs:
            r.font.size = Pt(round(new * (r.font.size.pt / size), 1))


def _fit_line(shape, text: str, lines: int = 1) -> None:
    """칸에 맞춘다: 한 줄에 들어가게 글자를 줄이고(최소 8pt), lines=2면 그래도 넘칠 때 7.4pt 두 줄로 쓴다.
    그래도 넘치면 괄호 내용을 빼고, 마지막으로 말줄임."""
    width = shape.width / 12700 - 16  # 좌우 기본 여백
    run = _runs(shape)[0]
    min_size = min(8.0, run.font.size.pt)
    if _text_width(text, min_size) > width and lines == 2:
        min_size, width = min(7.4, run.font.size.pt), width * 2 - 10  # 줄바꿈 손실 감안
    if _text_width(text, min_size) > width:
        text = PAREN.sub("", text).strip() or text
    while _text_width(text, min_size) > width and len(text) > 2:
        text = text[:-2] + "…"
    run.text = text
    _fit([run], text, width, min_size)


def _color(run, hex_: str) -> None:
    run.font.color.rgb = RGBColor.from_string(hex_)


def _rewrite_runs(paragraph, segments: list[tuple[str, str]], styles: dict) -> None:
    """문단의 run을 segments대로 다시 만든다. styles: 서식 이름 -> 본보기 run의 rPr XML."""
    for r in list(paragraph.runs):
        paragraph._p.remove(r._r)
    for text, style in segments:
        r = paragraph.add_run()  # endParaRPr 앞에 들어간다
        r._r.insert(0, copy.deepcopy(styles[style]))
        r.text = text


@dataclass
class Kpi:
    label: str
    value: str        # "-"이면 자료없음
    unit: str
    color: str        # NAVY | GREEN | ORANGE
    sub: str          # 괄호 안 보조 문구


@dataclass
class RadarAxis:
    label: str        # 성장성 등
    grade: str | None  # 우수·양호·보통·보통 이하·낮음 (없으면 None)
    score: int


@dataclass
class LetterData:
    """공문 한 장에 들어가는 기업별 값. PPTX(render_letter)와 모바일 이미지(mobile.py)가 같이 쓴다."""
    doc_no: str
    issued_text: str
    recipient_name: str
    recipient_suffix: str     # " 대표이사 홍길동 귀하"
    name: str                 # 짧은 기업명
    industry: str
    founded_text: str
    location: str
    basis_text: str           # KODATA · 2025년 결산 · …
    radar: list[RadarAxis]
    kpis: list[Kpi]
    checks: list[Check]
    counts: dict[str, int]
    verdict: list[tuple[str, str]]


def build_letter_data(c: Company, *, doc_no: str, issued: date | None = None) -> LetterData:
    issued = issued or date.today()
    addr = split_address(c.address)
    position = ((c.personal_info.get("name_position") or "").split("/") + [""])[1].strip()
    position = position or ("대표" if c.company_type == "개인사업자" else "대표이사")
    rep = (c.representative or "").strip()

    year, revenue = _latest(c.income_summary.get("매출액"))
    _, op = _latest(c.income_summary.get("영업이익"))
    settle = (c.settlement_date or "")[:4] or year
    founded = (c.founded_date or "")[:4]

    grades = c.diagnosis.model_dump() if c.diagnosis else {}
    radar = []
    for key, label in RADAR_AXES:
        g = norm_grade(grades.get(key))
        radar.append(RadarAxis(label, g, GRADE_SCORE.get(g, 0) if g else 0))

    prev_rev = _prev(c.income_summary.get("매출액"), year)
    prev_op = _prev(c.income_summary.get("영업이익"), year)
    rev_g = growth_rate(c, "매출액증가율", revenue, prev_rev)
    op_g = growth_rate(c, "영업이익증가율", op, prev_op)
    debt, debt_avg = indicator(c, "financial_structure", "부채비율")
    current, current_avg = indicator(c, "financial_structure", "유동비율")
    icr, _ = indicator(c, "debt_repayment", "이자보상배수(배)")
    _, equity = _latest(c.balance_summary.get("자본총계"))

    def paren_prev(prev, g):
        parts = [f"{fmt_eok(prev)}억" if prev is not None else None, fmt_signed_pct(g) if g is not None else None]
        parts = [p for p in parts if p]
        return f"({' · '.join(parts)})" if parts else "(전년 자료없음)"

    def trend_color(g, neg=False):
        if neg or (g is not None and g <= -10):
            return ORANGE
        return GREEN if g is not None and g >= 0 else NAVY

    rank = grade_rank(c.credit_grade)
    kpis = [
        Kpi("매출액", fmt_eok(revenue) if revenue is not None else "-", "억원", trend_color(rev_g), paren_prev(prev_rev, rev_g)),
        Kpi("영업이익", fmt_eok(op) if op is not None else "-", "억원", trend_color(op_g, neg=op is not None and op < 0),
            paren_prev(prev_op, op_g)),
        Kpi("부채비율", fmt_num(debt, 1) if debt is not None else "-", "%",
            NAVY if debt is None else GREEN if debt <= (debt_avg if debt_avg is not None else 200) else ORANGE,
            f"(업종 {fmt_num(debt_avg, 1)}%)" if debt_avg is not None else "(업종 자료없음)"),
        Kpi("유동비율", fmt_num(current, 1) if current is not None else "-", "%",
            NAVY if current is None else GREEN if current >= (current_avg if current_avg is not None else 100)
            else NAVY if current >= 100 else ORANGE,
            f"(업종 {fmt_num(current_avg, 1)}%)" if current_avg is not None else "(업종 자료없음)"),
        Kpi("신용등급", (c.credit_grade or "-").strip(), "",
            NAVY if rank is None else GREEN if rank <= grade_rank("bb-") else NAVY if rank < grade_rank("ccc+") else ORANGE,
            "(KODATA)"),
    ]
    checks = build_checks(c, addr, debt, debt_avg, current, icr, equity)
    return LetterData(
        doc_no=doc_no, issued_text=f"{issued.year}. {issued.month}. {issued.day}.",
        recipient_name=spaced_name(c.company_name),
        recipient_suffix=f" {position} {rep} 귀하" if rep else f" {position} 귀하",
        name=short_name(c.company_name), industry=pretty_industry(c.industry_name),
        founded_text=f"{founded}년 (업력 {issued.year - int(founded)}년)" if founded.isdigit() else "-",
        location=addr.short,
        basis_text=f"KODATA · {settle}년 결산 · 괄호: 전년 또는 업종평균" if settle else "KODATA · 괄호: 전년 또는 업종평균",
        radar=radar, kpis=kpis, checks=checks,
        counts={k: sum(ch.judgment == k for ch in checks) for k in JUDGE_STYLE},
        verdict=verdict_segments(checks),
    )


def render_letter(c: Company, *, doc_no: str, issued: date | None = None) -> bytes:
    d = build_letter_data(c, doc_no=doc_no, issued=issued)
    prs = Presentation(str(TEMPLATE))
    s1, s2 = prs.slides[0], prs.slides[1]

    # ----- 1쪽: 문서번호 · 시행일자 · 수신
    _set_text(_shape(s1, "Text 5"), d.doc_no)
    _set_text(_shape(s1, "Text 7"), d.issued_text)
    _set_text(_shape(s1, "Text 9"), d.recipient_name, d.recipient_suffix)

    # ----- 2쪽: 재무진단 요약
    _set_text(_shape(s2, "Text 7"), d.basis_text)
    _fit_line(_shape(s2, "Text 9"), d.name)
    _fit_line(_shape(s2, "Text 11"), d.industry, lines=2)
    _fit_line(_shape(s2, "Text 15"), d.location)
    _set_text(_shape(s2, "Text 13"), d.founded_text)

    # 오각형 그래프 (재무진단 5개 항목). 등급은 둘째 줄(괄호 중간에서 끊기지 않게), 맨 위 축(성장성)은 위 여백이 좁아 한 줄
    cats = [f"{a.label}({a.grade or '자료없음'})" if i == 0 else f"{a.label}\n({a.grade or '자료없음'})"
            for i, a in enumerate(d.radar)]
    chart_data = CategoryChartData()
    chart_data.categories = cats
    chart_data.add_series(d.name, [a.score for a in d.radar])
    chart_data.add_series("동종업종 평균", [INDUSTRY_AVG_SCORE] * len(cats))
    radar = _shape(s2, "radar")
    radar.chart.replace_data(chart_data)
    if not any(a.score for a in d.radar):
        box = s2.shapes.add_textbox(radar.left, radar.top + radar.height // 2 - Pt(10), radar.width, Pt(20))
        p = box.text_frame.paragraphs[0]
        p.alignment = 2  # 가운데
        run = p.add_run()
        run.text = "재무진단 자료 없음"
        run.font.size, run.font.bold = Pt(10), True
        _color(run, "6B7A8E")
    legend = _shape(s2, "Text 16")
    lr = _runs(legend)
    lr[1].text = d.name
    _fit([lr[0], lr[1]], "━ " + d.name, legend.width / 12700 - 6)
    # 등급 환산 안내(문단 3·4): 샘플 2줄 -> 우수·낮음을 더한 4줄
    paras = legend.text_frame.paragraphs
    for text in ("낮음 20", "보통 이하 40"):
        paras[4]._p.addnext(copy.deepcopy(paras[4]._p))
        legend.text_frame.paragraphs[5].runs[0].text = text
    paras = legend.text_frame.paragraphs
    paras[3].runs[0].text = "우수 90 · 양호 70"
    paras[4].runs[0].text = "보통 50"

    # KPI 5개 (값·단위 도형, 괄호 도형)
    kpi_shapes = [("Text 20", "Text 21"), ("Text 25", "Text 26"), ("Text 30", "Text 31"), ("Text 35", "Text 36"), ("Text 40", "Text 41")]
    for k, (value_shape, sub_shape) in zip(d.kpis, kpi_shapes):
        runs = _runs(_shape(s2, value_shape))
        runs[0].text = k.value
        if len(runs) > 1:
            runs[1].text = f" {k.unit}" if k.value != "-" and k.unit else ""
        for r in runs:
            _color(r, k.color)
        _set_text(_shape(s2, sub_shape), k.sub)

    # ----- 지원조건 점검 결과
    _set_text(_shape(s2, "Text 50"), f"✓ 충족 {d.counts['충족']}")
    _set_text(_shape(s2, "Text 48"), f"? 확인 필요 {d.counts['확인 필요']}")
    _set_text(_shape(s2, "Text 46"), f"! 보완 필요 {d.counts['보완 필요']}")
    table = _shape(s2, "Table 0").table
    for row, ch in zip(list(table.rows)[1:], d.checks):
        cells = row.cells
        cells[0].text_frame.paragraphs[0].runs[0].text = ch.label
        cells[1].text_frame.paragraphs[0].runs[0].text = ch.status
        fill, color, text = JUDGE_STYLE[ch.judgment]
        run = cells[2].text_frame.paragraphs[0].runs[0]
        run.text = text
        _color(run, color)
        cells[2].fill.solid()
        cells[2].fill.fore_color.rgb = RGBColor.from_string(fill)
        cells[3].text_frame.paragraphs[0].runs[0].text = ch.next_step

    # 종합 의견: 샘플 문단의 보통·굵게·초록 서식을 본보기로 다시 쓴다
    verdict = _shape(s2, "Text 54").text_frame.paragraphs[0]
    vr = verdict.runs
    styles = {"normal": vr[0]._r.rPr, "bold": vr[1]._r.rPr, "green": vr[3]._r.rPr}
    _rewrite_runs(verdict, d.verdict, styles)

    _shape(s2, "Text 94").text_frame.paragraphs[0].runs[0].text = f"{d.issued_text}   "

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def doc_number(issued: date, tier: str | None, seq: int) -> str:
    """FS-연도-월일-{등급 분류(S·A·B·C, 미지정 N)}{일련번호 2자리}. 샘플: FS-2026-1007-A02."""
    return f"FS-{issued.year}-{issued.month:02d}{issued.day:02d}-{tier or 'N'}{seq:02d}"
