"""그래프 데이터 표 → 엑셀. 화면의 모든 그래프 밑 표에 있는 '엑셀 다운로드'가 쓴다.
화면이 보여 준 표(제목·열·행)를 그대로 받아 xlsx로 돌려준다 — 그래프마다 엔드포인트를 따로 두지 않는다."""
from __future__ import annotations

import io
import re
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

router = APIRouter(prefix="/api", tags=["export"])

Cell = float | int | str | None


class TableExport(BaseModel):
    title: str = Field(min_length=1, max_length=100)
    subtitle: str | None = Field(default=None, max_length=200)
    columns: list[str] = Field(min_length=1, max_length=30)
    rows: list[list[Cell]] = Field(max_length=5000)


_UNSAFE_FILENAME = re.compile(r'[\\/:*?"<>|]')
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_table_xlsx(data: TableExport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", data.title)[:31]
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([data.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([data.subtitle or f"다운로드 {datetime.now():%Y-%m-%d %H:%M}"])
    ws["A2"].font = Font(color="64748B", size=10)
    ws.append([])

    ws.append(data.columns)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F5F9")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row in data.rows:
        ws.append(list(row) + [None] * (len(data.columns) - len(row)))
        for cell in ws[ws.max_row]:
            cell.border = border
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0" if float(cell.value).is_integer() else "#,##0.0#"

    for i, name in enumerate(data.columns, start=1):
        values = [str(name)] + [str(r[i - 1]) for r in data.rows if i - 1 < len(r) and r[i - 1] is not None]
        width = max(len(v.encode("utf-8")) * 0.75 for v in values) + 4
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/export/table")
def export_table(data: TableExport):
    filename = f"{_UNSAFE_FILENAME.sub('_', data.title)}_{datetime.now():%Y%m%d}.xlsx"
    return Response(build_table_xlsx(data), media_type=XLSX,
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
